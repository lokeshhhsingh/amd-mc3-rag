#!/usr/bin/env python3
"""Builds a harder practice corpus + questions for testing the MC3 solution.

    python3 build_hard_corpus.py            # writes mc3-corpus-hard/ and hard-questions.json next to this file

Fictional company "Helios Robotics" (HX-9, HX-9L, HX-12 robot arms). The traps mirror the README's:
withdrawn / draft / superseded documents, near-duplicate part numbers and products, big noisy logs and
tables, facts that exist only inside images (with uninformative file names), multi-hop questions
(log -> error KB -> price list), unanswerable questions, and unreadable files.
Everything is generated from a fixed seed, so the expected answers below are exact.
"""
from __future__ import annotations

import csv
import json
import os
import random
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import docx
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE / "mc3-corpus-hard"
random.seed(20261018)
MONO_B = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf"
SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def w(rel: str, text: str) -> None:
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")


def make_pdf(rel: str, pages: list, password: str | None = None) -> None:
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(p), pagesize=letter)
    for lines in pages:
        y = 740
        for i, ln in enumerate(lines):
            c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 14 if i == 0 else 11)
            c.drawString(60, y, ln)
            y -= 24 if i == 0 else 18
        c.showPage()
    c.save()
    if password:
        r, wr = PdfReader(str(p)), PdfWriter()
        for pg in r.pages:
            wr.add_page(pg)
        wr.encrypt(password, password + "-owner")
        with open(p, "wb") as fh:
            wr.write(fh)


def spec_pdf(rel, title, sub, rows, extra=None):
    page1 = [title, sub, "", "Helios Robotics GmbH, Industrial Arms Division", "",
             "This document describes the mechanical and electrical data of the arm.",
             "Operating temperature 0 to 45 C. Storage temperature -20 to 60 C."] + (extra or [])
    page2 = ["Technical data"] + [f"{k} {'.' * max(3, 46 - len(k))} {v}" for k, v in rows]
    make_pdf(rel, [page1, page2])


# ------------------------------------------------------------------ specs (PDF)
spec_pdf("engineering/specs/hx9_datasheet_rev1_WITHDRAWN.pdf", "HX-9 datasheet, revision 1",
         "WITHDRAWN - superseded by revision 2. Do not use.",
         [("Payload", "9.0 kg"), ("Reach", "900 mm"), ("Repeatability", "+/-0.05 mm"), ("Protection class", "IP54"),
          ("Maximum joint speed", "210 deg/s"), ("Power consumption", "380 W")])
spec_pdf("engineering/specs/hx9_datasheet_rev2.pdf", "HX-9 datasheet, revision 2",
         "Supersedes revision 1. Current release.",
         [("Payload", "8.5 kg"), ("Reach", "900 mm"), ("Repeatability", "+/-0.03 mm"), ("Protection class", "IP65"),
          ("Maximum joint speed", "230 deg/s"), ("Power consumption", "400 W")],
         extra=["Revision history: revision 1 was withdrawn after a payload test."])
spec_pdf("engineering/specs/hx12_datasheet_rev3.pdf", "HX-12 datasheet, revision 3", "Current release.",
         [("Payload", "12.0 kg"), ("Reach", "1250 mm"), ("Repeatability", "+/-0.04 mm"), ("Protection class", "IP67"),
          ("Maximum joint speed", "180 deg/s"), ("Power consumption", "650 W")])
make_pdf("engineering/specs/hx9L_datasheet_DRAFT.pdf",
         [["HX-9L datasheet (DRAFT)", "DRAFT - not approved for release. Targets only.", "",
           "Payload target ........ 7.0 kg", "Reach target ........ 950 mm"]])

d = docx.Document()
d.add_heading("HX-9L product brief", 0)
d.add_paragraph("Status: approved for the pilot programme. This brief replaces the HX-9L draft datasheet.")
d.add_heading("Approved data", 1)
t = d.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Parameter", "Value"
for k, v in [("Payload", "7.5 kg"), ("Reach", "1000 mm"), ("Protection class", "IP54")]:
    r = t.add_row().cells
    r[0].text, r[1].text = k, v
d.add_paragraph("Repeatability has not been specified yet.")
p = ROOT / "engineering/specs/hx9L_product_brief.docx"
p.parent.mkdir(parents=True, exist_ok=True)
d.save(p)

# ------------------------------------------------------------------ controller manual (DOCX with footer)
d = docx.Document()
d.add_heading("Helios controller manual", 0)
d.add_heading("1. Network", 1)
d.add_paragraph("The controller listens on port 7420 by default. The service port 7421 is used only by field technicians.")
d.add_heading("2. Security", 1)
d.add_paragraph("Admin sessions expire after 15 minutes of inactivity. Operator sessions expire after 60 minutes.")
d.add_heading("3. Error codes", 1)
t = d.add_table(rows=1, cols=3)
for c, txt in zip(t.rows[0].cells, ["Code", "Meaning", "Action"]):
    c.text = txt
for row in [("E101", "Encoder fault", "Recalibrate"), ("E102", "Over-temperature", "Let the arm cool"),
            ("E205", "Brake not released", "Check the brake relay")]:
    cells = t.add_row().cells
    for c, txt in zip(cells, row):
        c.text = txt
d.sections[0].footer.paragraphs[0].text = "Document owner: Priya Nair | Controller manual v5.2"
p = ROOT / "engineering/manuals/controller_manual.docx"
p.parent.mkdir(parents=True, exist_ok=True)
d.save(p)

# ------------------------------------------------------------------ firmware, code, config, KB
w("engineering/firmware/release_notes.txt", """Helios controller firmware release notes

Firmware 5.1.0 (2026-05-02)
  - New admin session timeout setting.
  - Fixes HX-ERR-4388 (spurious over-temperature alarm).

Firmware 5.2.0 (2026-06-18)
  - Fixes HX-ERR-4410 (encoder drift on joint J2).
  - Fixes HX-ERR-4412 (servo bus timeout under load).

Firmware 5.2.1 (2026-07-30)
  - Fixes HX-ERR-4471 (intermittent brake release).
  - Minor logging improvements.
""")
w("engineering/firmware/controller/config.py", '''"""Controller defaults."""

# Raised from 30 in firmware 5.1.0.
DEFAULT_TIMEOUT_S = 45
MAX_RETRIES = 4
HEARTBEAT_INTERVAL_S = 20
''')
w("engineering/firmware/controller/config_legacy.py", '''"""Legacy defaults, kept for the HX-6 line only."""
DEFAULT_TIMEOUT_S = 30
MAX_RETRIES = 3
''')
w("engineering/firmware/deploy/settings.yaml", "retry_limit: 6\nbroker_port: 5671\nheartbeat_s: 20\nlog_level: info\n")
w("engineering/kb/error_codes.md", """# Helios controller error codes

## HX-ERR-4388
Spurious over-temperature alarm. Check the thermistor cable.

## HX-ERR-4410
Encoder drift on joint J2. Recalibrate the encoder.

## HX-ERR-4412
Servo bus timeout. Check cable X3 and the bus terminator.

## HX-ERR-4471
Brake release failure caused by a worn relay. Replace relay part BRK-REL-77 and clear the fault.

## HX-ERR-4502
Safety loop open. Check the emergency-stop chain.
""")

# ------------------------------------------------------------------ logs (big, noisy)
def noisy_log(env: str, lines: int, inject: dict, fw: str) -> str:
    msgs = ["cycle complete", "heartbeat ok", "tool change", "queue depth 3", "temperature nominal", "calibration check passed",
            "conveyor sync", "gripper closed", "gripper open", "payload verified"]
    out = []
    for i in range(lines):
        day = 1 + (i * 28) // lines
        sec = random.randint(0, 86399)
        ts = f"2026-08-{day:02d} {sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"
        out.append((ts, f"{ts} [INFO] line{random.randint(1, 6)} {random.choice(msgs)} fw={fw} cell=C{random.randint(1, 6)}-0{random.randint(1, 4)}"))
    for ts, text in inject.items():
        out.append((ts, f"{ts} {text}"))
    out.sort()
    return "\n".join(t for _, t in out) + "\n"

w("ops/logs/production_line4_2026-08.log", noisy_log("prod", 4000, {
    "2026-08-09 14:20:11": "[ERROR] line2 STOP code=HX-ERR-4410 fw=5.2.0 cell=C2-01",
    "2026-08-14 03:12:09": "[ERROR] line4 STOP code=HX-ERR-4471 fw=5.2.0 cell=C4-02",
    "2026-08-14 03:12:10": "[WARN] line4 brake release timeout, operator notified",
    "2026-08-21 22:05:40": "[ERROR] line3 STOP code=HX-ERR-4502 fw=5.2.0 cell=C3-03"}, "5.2.0"))
w("ops/logs/staging_line4_2026-08.log", noisy_log("staging", 600, {
    "2026-08-14 02:58:41": "[ERROR] line4 STOP code=HX-ERR-4412 fw=5.1.0 cell=C4-01"}, "5.1.0"))
w("ops/logs/startup_2026-09-01.log", "2026-09-01 06:00:01 [INFO] controller boot fw=5.2.1\n2026-09-01 06:00:09 [INFO] self test passed\n")

# ------------------------------------------------------------------ tickets (CSV), RMA + prices (XLSX/CSV)
comps = ["encoder", "brake relay", "gripper", "pendant", "power supply", "cable", "firmware", "fan"]
with open(ROOT / "support/tickets/ticket_export.csv", "w", newline="") if (ROOT / "support/tickets").mkdir(parents=True, exist_ok=True) is None else None as fh:
    wr = csv.writer(fh)
    wr.writerow(["ticket_id", "product", "component", "severity", "status", "fixed_in", "summary"])
    for n in range(1000, 1300):
        if n == 1187:
            wr.writerow(["HRT-1187", "HX-12", "brake relay", "S2", "Closed", "5.2.1", "Intermittent brake release, relay wear"])
        elif n == 1190:
            wr.writerow(["HRT-1190", "HX-9", "encoder", "S3", "Closed", "5.2.0", "Encoder drift on joint J2"])
        else:
            wr.writerow([f"HRT-{n}", random.choice(["HX-9", "HX-12", "HX-9L"]), random.choice(comps), random.choice(["S1", "S2", "S3"]),
                         random.choice(["Open", "Closed", "Closed"]), random.choice(["5.1.0", "5.2.0", "5.2.1", ""]), "Routine service request"])

def workbook(rel, sheets):
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        if name in ("USD", "EUR"):            # numbers stored as 38.5 but DISPLAYED as 38.50, like a real price list
            for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
                row[0].number_format = "0.00"
    wb.save(p)

parts = [(f"SEN-{random.randint(1000, 9999)}", f"Sensor type {i}", round(random.uniform(5, 90), 2)) for i in range(60)]
def prices(special, scale):
    return [["part_number", "description", "unit_price"]] + [["BRK-REL-77", "Brake release relay", special[0]],
            ["BRK-REL-77B", "Brake release relay, high-duty", special[1]]] + [[a, b, round(c * scale, 2)] for a, b, c in parts]
workbook("finance/price_list_2026.xlsx", {"USD": prices((38.50, 41.20), 1.0), "EUR": prices((35.90, 38.40), 0.93)})
with open(ROOT / "finance/price_list_2025.csv", "w", newline="") as fh:
    wr = csv.writer(fh)
    wr.writerow(["part_number", "description", "unit_price_usd"])
    wr.writerow(["BRK-REL-77", "Brake release relay", "34.00"])
    for a, b, c in parts[:20]:
        wr.writerow([a, b, round(c * 0.95, 2)])
workbook("support/rma/rma_log.xlsx", {
    "2026-Q2": [["rma_id", "part_number", "qty", "received", "disposition"], ["RMA-0412", "BRK-REL-77", 5, "2026-05-20", "scrapped"],
                ["RMA-0420", "SEN-2210", 2, "2026-06-11", "repaired"]],
    "2026-Q3": [["rma_id", "part_number", "qty", "received", "disposition"], ["RMA-0501", "BRK-REL-77", 12, "2026-08-18", "scrapped"],
                ["RMA-0502", "BRK-REL-77B", 3, "2026-08-19", "scrapped"], ["RMA-0517", "SEN-4471", 1, "2026-09-02", "repaired"]],
    "Notes": [["Returns are counted when received, not when shipped."]]})

# ------------------------------------------------------------------ policies, meetings, contracts, audit
make_pdf("finance/contracts/vorlon_supplier_agreement_ENCRYPTED.pdf",
         [["Vorlon supplier agreement", "Payment terms: net 45 days from invoice date.", "Governing law: Germany."]], password="s3cret-pw")
w("finance/audit/internal_audit_2026.txt", "Internal audit 2026\nNumber of findings: 7\n")
d = docx.Document()
d.add_heading("Travel policy", 0)
t = d.add_table(rows=1, cols=3)
for c, txt in zip(t.rows[0].cells, ["Category", "Domestic", "International"]):
    c.text = txt
for row in [("Daily meal allowance", "45", "70"), ("Hotel cap per night", "140", "220")]:
    for c, txt in zip(t.add_row().cells, row):
        c.text = txt
p = ROOT / "hr/travel_policy.docx"
p.parent.mkdir(parents=True, exist_ok=True)
d.save(p)
make_pdf("hr/leave_policy.pdf", [["Leave policy", "Annual leave entitlement: 24 working days.", "Maximum carry-over: 5 days.",
                                  "Carried-over days must be used by 31 March."]])
w("meetings/2026-08-20_pilot_planning.txt", "Pilot planning meeting, 20 August 2026\nKestrel Foods pilot: tentative ship date 2026-10-30.\n")
w("meetings/2026-09-03_pilot_decision.txt", "Pilot decision meeting, 3 September 2026\nDecision: the Kestrel Foods pilot ship date moves to 2026-11-15, replacing the 2026-10-30 estimate.\n")

# ------------------------------------------------------------------ images (uninformative names, degraded)
def text_image(rel, lines, size=(900, 520), bg=(235, 235, 235), fg=(20, 20, 20), rot=0.0, blur=0.0, dim=1.0, boxes=None):
    im = Image.new("RGB", size, bg)
    dr = ImageDraw.Draw(im)
    f = ImageFont.truetype(MONO_B, 30)
    y = 40
    for ln in lines:
        dr.text((50, y), ln, font=f, fill=fg)
        y += 52
    for b in boxes or []:
        dr.rectangle(b, outline=fg, width=3)
    if rot:
        im = im.rotate(rot, expand=True, fillcolor=bg, resample=Image.BICUBIC)
    if blur:
        im = im.filter(ImageFilter.GaussianBlur(blur))
    if dim != 1.0:
        im = ImageEnhance.Brightness(im).enhance(dim)
    p = ROOT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    im.save(p, quality=80) if rel.endswith("jpg") else im.save(p)

text_image("media/IMG_0412.jpg", ["HELIOS ROBOTICS", "MODEL HX-12", "SERIAL HR-2207-0418", "MFG DATE 2026-03", "MADE IN GERMANY"],
           bg=(150, 155, 160), fg=(30, 30, 30), rot=2.0, blur=0.8, dim=0.8)
text_image("media/scan_0007.png", ["JOINT   TORQUE (N-m)", "J1      60", "J2      55", "J3      42", "J4      28", "J5      18", "J6      9"],
           size=(900, 560), rot=-1.5, blur=0.5)
text_image("media/diagram_final2.png", ["CONNECTOR X7", "pin 12  24V_IN", "pin 13  GND", "pin 14  BRAKE_REL", "pin 15  ESTOP_A", "pin 16  ESTOP_B"],
           size=(900, 520), bg=(250, 250, 250), boxes=[(30, 20, 700, 330)], rot=1.0)

# ------------------------------------------------------------------ junk, traps, edge cases
(ROOT / "misc").mkdir(parents=True, exist_ok=True)
(ROOT / "misc/telemetry.bin").write_bytes(bytes(random.randrange(256) for _ in range(4096)))
(ROOT / "misc/zero.txt").write_bytes(b"")
(ROOT / "misc/corrupt.pdf").write_bytes(b"%PDF-1.4\nthis is not really a pdf\n" + bytes(random.randrange(256) for _ in range(300)))
(ROOT / "misc/empty_folder").mkdir(exist_ok=True)
w("misc/Uebersicht Notizen.txt", "Kantinenplan: Montag Suppe, Dienstag Pasta.\n")
w("archive/2024/q4/old/legacy/readme.txt", "Legacy HX-6 notes. Default timeout was 30 seconds.\n")

# ------------------------------------------------------------------ questions
Q = [
 ("What is the payload of the HX-9 in kg?", "8.5", ["engineering/specs/hx9_datasheet_rev2.pdf"], "withdrawn r1 says 9.0"),
 ("What is the repeatability of the HX-12 in mm?", "0.04", ["engineering/specs/hx12_datasheet_rev3.pdf"], ""),
 ("What protection class does the current HX-9 datasheet give?", "IP65", ["engineering/specs/hx9_datasheet_rev2.pdf"], "withdrawn r1 says IP54"),
 ("What is the approved payload of the HX-9L in kg?", "7.5", ["engineering/specs/hx9L_product_brief.docx"], "DRAFT datasheet says 7.0"),
 ("What is the reach of the HX-9L in mm?", "1000", ["engineering/specs/hx9L_product_brief.docx"], "DRAFT says 950"),
 ("What is the default network port of the controller?", "7420", ["engineering/manuals/controller_manual.docx"], "7421 is the service port"),
 ("Who is the document owner of the controller manual?", "Priya Nair", ["engineering/manuals/controller_manual.docx"], "only in the footer"),
 ("After how many minutes do admin sessions expire?", "15", ["engineering/manuals/controller_manual.docx"], "operator is 60"),
 ("What does error E101 mean?", "Encoder fault", ["engineering/manuals/controller_manual.docx"], "table lookup"),
 ("Which firmware release fixed HX-ERR-4471?", "5.2.1", ["engineering/firmware/release_notes.txt"], ""),
 ("Which firmware release fixed the error that stopped production line 4 on 2026-08-14?", "5.2.1",
  ["engineering/firmware/release_notes.txt", "ops/logs/production_line4_2026-08.log"], "multi-hop: log -> release notes; staging log has a decoy"),
 ("What is the USD unit price of the part that must be replaced for the error that stopped production line 4 on 2026-08-14?", "38.50",
  ["engineering/kb/error_codes.md", "finance/price_list_2026.xlsx", "ops/logs/production_line4_2026-08.log"], "TWO hops: log -> KB -> price list"),
 ("What is the 2026 USD unit price of BRK-REL-77?", "38.50", ["finance/price_list_2026.xlsx"], "2025 csv and EUR sheet are decoys"),
 ("What is the 2026 USD unit price of BRK-REL-77B?", "41.20", ["finance/price_list_2026.xlsx"], "near-duplicate part number"),
 ("What was the USD unit price of BRK-REL-77 in the 2025 price list?", "34.00", ["finance/price_list_2025.csv"], ""),
 ("How many units of BRK-REL-77 were returned in 2026-Q3?", "12", ["support/rma/rma_log.xlsx"], "Q2 row says 5"),
 ("What torque is specified for joint J4?", "28", ["media/scan_0007.png"], "image only, file name uninformative"),
 ("What serial number is printed on the HX-12 nameplate?", "HR-2207-0418", ["media/IMG_0412.jpg"], "image only, dim and rotated"),
 ("Which pin of connector X7 carries BRAKE_REL?", "14", ["media/diagram_final2.png"], "image only"),
 ("What is DEFAULT_TIMEOUT_S in config.py?", "45", ["engineering/firmware/controller/config.py"], "config_legacy.py says 30"),
 ("What is MAX_RETRIES in config.py?", "4", ["engineering/firmware/controller/config.py"], ""),
 ("What broker port is set in settings.yaml?", "5671", ["engineering/firmware/deploy/settings.yaml"], ""),
 ("What is the current planned ship date of the Kestrel Foods pilot?", "2026-11-15", ["meetings/2026-09-03_pilot_decision.txt"], "later meeting overrides the earlier estimate"),
 ("What is the daily meal allowance for domestic travel?", "45", ["hr/travel_policy.docx"], "international is 70"),
 ("What is the maximum number of leave days that can be carried over?", "5", ["hr/leave_policy.pdf"], ""),
 ("In which firmware release was ticket HRT-1187 fixed?", "5.2.1", ["support/tickets/ticket_export.csv"], "one row among 300"),
 ("Which firmware version was line 4 running at the production stop on 2026-08-14?", "5.2.0", ["ops/logs/production_line4_2026-08.log"], "one line among 4000"),
 ("What is the unit price of BRK-REL-77 in GBP?", "", [], "UNANSWERABLE: only USD and EUR exist"),
 ("What is the maximum payload of the HX-15?", "", [], "UNANSWERABLE: no such product"),
 ("What is the repeatability of the HX-9L?", "", [], "UNANSWERABLE: brief says not specified; draft has no value"),
 ("What payment terms does the Vorlon supplier agreement specify?", "", [], "answer exists only in the ENCRYPTED pdf: must refuse"),
 ("How many findings did the 2026 internal audit report?", "", [], "answer exists only in an UNREADABLE file (needs a non-root user to test)"),
]
queries = []
for n, (q, a, cites, note) in enumerate(Q, start=1):
    item = {"n": n, "query": q, "expected_answer": a, "expected_citations": cites, "note": note}
    if "UNREADABLE" in note:
        item["skip_if_root"] = True
    queries.append(item)
(HERE / "hard-questions.json").write_text(json.dumps({"challenge": "hard practice set", "queries": queries}, indent=1))

# the one unreadable file: only effective for non-root users
os.chmod(ROOT / "finance/audit/internal_audit_2026.txt", 0o000)
print("built", sum(1 for _ in ROOT.rglob("*") if _.is_file()), "files and", len(queries), "questions")
